"""Anexo A del trabajo integrador: cálculos de pronóstico, modelo y factibilidad en Excel.

Lo llama el notebook 03 (sección 6). Las hojas de flujo replican, con fórmulas de Excel,
las mismas reglas de backend/observatorio/factibilidad.py: si se cambia un supuesto en la
hoja «Supuestos», todo se recalcula. El notebook verifica después que Excel y Python den lo
mismo.

Convención de colores (modelos financieros): azul = dato o supuesto que se puede cambiar;
negro = fórmula; verde = vínculo a otra hoja; relleno amarillo = supuestos clave.
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from observatorio import factibilidad as fac

FUENTE = 'Arial'
AZUL, NEGRO, VERDE = '0000FF', '000000', '008000'
AMARILLO = PatternFill('solid', fgColor='FFFF00')
ENCABEZADO = PatternFill('solid', fgColor='E8E6FD')
GRIS = PatternFill('solid', fgColor='F2F2F2')
BORDE = Border(bottom=Side(style='thin', color='BBBBBB'))
PESOS = '$#,##0;($#,##0);"-"'
PESOS2 = '$#,##0.00;($#,##0.00);"-"'
PCT = '0.0%;(0.0%);"-"'
NUM1 = '#,##0.0;(#,##0.0);"-"'
NUM0 = '#,##0;(#,##0);"-"'
NUM3 = '0.000;(0.000);"-"'

ESC_COL = {'Pesimista': 'C', 'Realista': 'D', 'Optimista': 'E'}

RIESGOS = [
    ('Mercado', 'Menor adopción de la prevista: solo 3.2 % de las cuidadoras buscó un servicio formal y 32.2 % no '
                'contrataría ninguno (ENASIC 2022, nacional); la desconfianza en alguien ajeno es una razón citada.',
     'Media', 'Alto', 'Cuidadoras de las mismas colonias, verificación visible y piloto de 15 familias; revisión en '
                      'el mes 12 con regla de cierre (≥ 50 % de las familias del punto de equilibrio).'),
    ('Mercado', 'Los 300 Centros de Cuidado Infantil gratuitos anunciados por la CDMX (sep-2026) pueden reducir '
                'la demanda del segmento de 0 a 5 años.',
     'Media', 'Medio', 'Enfocar la red en personas mayores y con discapacidad y en horarios fuera de 6:30 a 18:00; '
                       'presentarla como complemento del Sistema Público de Cuidados, no como competencia.'),
    ('Costos', 'Aumentos del salario mínimo por encima de la tarifa (el general subió 13 % en 2026) comprimen el '
               'margen de $25.50 por hora.',
     'Alta', 'Medio', 'Indexar el copago y el subsidio al salario mínimo profesional del trabajo del hogar dentro '
                      'del convenio.'),
    ('Ejecución', 'Rotación de cuidadoras e incidentes de seguridad o maltrato en el domicilio.',
     'Media', 'Alto', 'Curso de 40 h, verificación de antecedentes, seguro de responsabilidad civil, protocolo de '
                      'incidentes y una trabajadora social por cada 80 familias.'),
    ('Financiero', 'Dependencia del subsidio: con menos de ~$38 por hora el VAN realista es negativo.',
     'Media', 'Alto', 'Firmar un convenio plurianual antes de la etapa 2 y buscar fuentes complementarias '
                      '(donativos, responsabilidad social empresarial).'),
    ('Legal', 'Tratamiento de datos sensibles (salud, discapacidad, ubicación) de familias y cuidadoras bajo la '
              'LFPDPPP vigente desde el 21 de marzo de 2025.',
     'Baja', 'Alto', 'Aviso de privacidad, minimización de datos, seudonimización en los indicadores y protocolo '
                     'de gobernanza (entregables de la UCA de Leyes).'),
]


def _f(color=NEGRO, negrita=False, tam=10, italica=False):
    return Font(name=FUENTE, color=color, bold=negrita, size=tam, italic=italica)


def _celda(ws, ref, valor, color=NEGRO, formato=None, negrita=False, relleno=None, nota=None, alinear=None):
    c = ws[ref]
    c.value = valor
    c.font = _f(color, negrita)
    if formato:
        c.number_format = formato
    if relleno:
        c.fill = relleno
    if nota:
        c.comment = Comment(nota, 'PP7')
    if alinear:
        c.alignment = alinear
    return c


def _titulo(ws, texto, sub=None):
    ws['A1'].value = texto
    ws['A1'].font = _f(negrita=True, tam=13)
    if sub:
        ws['A2'].value = sub
        ws['A2'].font = _f(italica=True, tam=9)


def _encabezados(ws, fila, textos, col_inicial=1):
    for i, t in enumerate(textos):
        c = ws.cell(row=fila, column=col_inicial + i, value=t)
        c.font = _f(negrita=True)
        c.fill = ENCABEZADO
        c.alignment = Alignment(wrap_text=True, vertical='center')


def _anchos(ws, anchos):
    for col, ancho in anchos.items():
        ws.column_dimensions[col].width = ancho


# =========================================================================================
# Hoja «Supuestos»
# =========================================================================================

COMUNES = [
    ('hogares_zona', 'Hogares con necesidad de cuidado en la zona', 'hogares', NUM0),
    ('copago_hora', 'Copago de la familia', '$ por hora', PESOS2),
    ('pago_cuidadora_hora', 'Pago bruto a la cuidadora', '$ por hora', PESOS2),
    ('prestaciones', 'Prestaciones y seguridad social sobre el pago', '%', PCT),
    ('insumos_hora', 'Seguro por hora, insumos y traslados', '$ por hora', PESOS2),
    ('horas_semana_cuidadora', 'Jornada de una cuidadora de la red', 'horas por semana', NUM1),
    ('rotacion_mensual', 'Rotación mensual de cuidadoras', '%', PCT),
    ('capacitacion_cuidadora', 'Capacitación y verificación por cuidadora nueva', '$', PESOS),
    ('familias_por_trabajadora_social', 'Familias por trabajadora social', 'familias', NUM0),
    ('familias_piloto', 'Familias del piloto', 'familias', NUM0),
    ('cuidadoras_piloto', 'Cuidadoras del piloto', 'cuidadoras', NUM0),
    ('meses_piloto', 'Duración del piloto', 'meses', NUM0),
    ('coordinacion', 'Coordinadora del programa (con 35 % de carga social)', '$ al mes', PESOS),
    ('trabajadora_social', 'Trabajadora social (con 35 % de carga social)', '$ al mes', PESOS),
    ('plataforma', 'Plataforma: hospedaje, mensajería y mantenimiento', '$ al mes', PESOS),
    ('seguro_fijo', 'Póliza de responsabilidad civil', '$ al mes', PESOS),
    ('difusion', 'Trabajo comunitario y difusión', '$ al mes', PESOS),
    ('administracion', 'Contabilidad y facturación', '$ al mes', PESOS),
    ('inv_plataforma', 'Inversión: adaptar la plataforma de vinculación', '$', PESOS),
    ('inv_proteccion_datos', 'Inversión: protección de datos (LFPDPPP)', '$', PESOS),
    ('inv_equipamiento', 'Inversión: equipamiento (kits, botiquines, teléfonos)', '$', PESOS),
    ('inv_lanzamiento', 'Inversión: lanzamiento en las colonias', '$', PESOS),
    ('tasa_descuento_anual', 'Tasa de descuento real anual', '%', PCT),
    ('colchon', 'Colchón para imprevistos', '%', PCT),
    ('mes_revision', 'Mes de revisión (decisión por etapas)', 'mes', NUM0),
    ('umbral_continuar', 'Umbral para continuar (fracción de familias del punto de equilibrio)', '%', PCT),
    ('meses_costo_cierre', 'Costo de cierre, en meses de costos fijos', 'meses', NUM1),
]
NOTAS_COMUNES = {
    'coordinacion': '$22,000 brutos al mes × 1.35.',
    'trabajadora_social': '$18,000 brutos al mes × 1.35. Se agrega una por cada bloque de 80 familias.',
    'inv_plataforma': 'Valor de mercado de adaptar el prototipo del equipo (seguridad, perfiles, bitácora).',
    'capacitacion_cuidadora': 'Curso de 40 horas, constancia de antecedentes y kit inicial.',
    'mes_revision': 'Si en este mes hay menos familias que el umbral, el programa se cierra.',
}


def _hoja_supuestos(wb, hogares_zona):
    ws = wb.create_sheet('Supuestos')
    _titulo(ws, 'Supuestos del modelo financiero — Red de Cuidado a Domicilio «Manos que Cuidan» (escenario simulado)',
            'Pesos constantes de 2026. Celdas en azul: supuestos que se pueden cambiar. Relleno amarillo: variables clave por escenario.')
    _encabezados(ws, 4, ['Variable que cambia por escenario', 'Unidad', 'Pesimista', 'Realista', 'Optimista', 'Fuente o justificación'])
    refs = {}
    fila = 5
    formatos = {'adopcion': PCT, 'horas_semana_familia': NUM1, 'subsidio_hora': PESOS2, 'meses_estado_estable': NUM0}
    for campo, etiqueta, unidad in fac.VARIABLES_ESCENARIO:
        _celda(ws, f'A{fila}', etiqueta)
        _celda(ws, f'B{fila}', unidad)
        for nombre, col in ESC_COL.items():
            _celda(ws, f'{col}{fila}', getattr(fac.ESCENARIOS[nombre], campo), AZUL, formatos[campo], relleno=AMARILLO)
        _celda(ws, f'F{fila}', fac.SUPUESTOS_FUENTES.get(campo, ''))
        refs[campo] = fila
        fila += 1
    _celda(ws, f'A{fila}', 'Probabilidad asignada')
    _celda(ws, f'B{fila}', '%')
    for nombre, col in ESC_COL.items():
        _celda(ws, f'{col}{fila}', fac.PROBABILIDADES[nombre], AZUL, PCT, relleno=AMARILLO)
    _celda(ws, f'F{fila}', 'Subjetiva y cargada hacia abajo: no hay historia del servicio y el subsidio depende de un convenio que no existe.')
    refs['probabilidad'] = fila
    fila += 1
    _celda(ws, f'A{fila}', 'Suma de probabilidades (debe ser 100 %)')
    _celda(ws, f'C{fila}', f'=SUM(C{fila - 1}:E{fila - 1})', formato=PCT)
    fila += 2

    _encabezados(ws, fila, ['Parámetro común', 'Unidad', 'Valor', '', '', 'Fuente o justificación'])
    fila += 1
    base = fac.ESCENARIOS['Realista']
    for campo, etiqueta, unidad, formato in COMUNES:
        valor = hogares_zona if campo == 'hogares_zona' else getattr(base, campo)
        _celda(ws, f'A{fila}', etiqueta)
        _celda(ws, f'B{fila}', unidad)
        _celda(ws, f'C{fila}', valor, AZUL, formato, nota=NOTAS_COMUNES.get(campo))
        _celda(ws, f'F{fila}', fac.SUPUESTOS_FUENTES.get(campo, ''))
        refs[campo] = fila
        fila += 1

    fila += 1
    _encabezados(ws, fila, ['Derivado (fórmula)', 'Unidad', 'Valor', '', '', 'Cómo se calcula'])
    fila += 1
    c = lambda k: f'$C${refs[k]}'
    derivados = [
        ('semanas_mes', 'Semanas por mes', 'semanas', '=52/12', NUM3, '52 semanas ÷ 12 meses'),
        ('costo_variable_hora', 'Costo variable por hora', '$ por hora',
         f'={c("pago_cuidadora_hora")}*(1+{c("prestaciones")})+{c("insumos_hora")}', PESOS2, 'Pago × (1 + prestaciones) + insumos'),
        ('fijos_base', 'Costos fijos sin trabajadoras sociales', '$ al mes',
         f'={c("coordinacion")}+{c("plataforma")}+{c("seguro_fijo")}+{c("difusion")}+{c("administracion")}', PESOS,
         'Coordinación + plataforma + seguro + difusión + administración'),
        ('inversion_inicial', 'Inversión inicial (mes 0)', '$',
         f'={c("inv_plataforma")}+{c("inv_proteccion_datos")}+{c("inv_equipamiento")}+{c("inv_lanzamiento")}'
         f'+{c("cuidadoras_piloto")}*{c("capacitacion_cuidadora")}', PESOS, 'Inversiones + capacitación de las cuidadoras del piloto'),
        ('tasa_mensual', 'Tasa de descuento mensual equivalente', '%', f'=(1+{c("tasa_descuento_anual")})^(1/12)-1', '0.000%',
         '(1 + tasa anual)^(1/12) − 1'),
    ]
    for clave, etiqueta, unidad, formula, formato, como in derivados:
        _celda(ws, f'A{fila}', etiqueta)
        _celda(ws, f'B{fila}', unidad)
        _celda(ws, f'C{fila}', formula, formato=formato)
        _celda(ws, f'F{fila}', como)
        refs[clave] = fila
        fila += 1
    _celda(ws, f'A{fila + 1}', 'El horizonte es de 60 meses (vida útil de la plataforma y del equipo), sin valor de rescate. '
                               'Cambiarlo requiere agregar o quitar filas en las hojas de flujo.', alinear=Alignment(wrap_text=False))
    ws[f'A{fila + 1}'].font = _f(italica=True, tam=9)
    _anchos(ws, {'A': 58, 'B': 18, 'C': 14, 'D': 14, 'E': 14, 'F': 110})
    ws.freeze_panes = 'A5'
    return refs


# =========================================================================================
# Hojas de flujo
# =========================================================================================

PRIMERA = 6
COLUMNAS_FLUJO = ['Mes', 'Familias', 'Horas por semana', 'Horas al mes', 'Cuidadoras', 'Cuidadoras nuevas (incluye rotación)',
                  'Trabajadoras sociales', 'Ingresos ($)', 'Costo variable ($)', 'Costos fijos ($)', 'Capacitación ($)',
                  'Inversión ($)', 'Flujo neto ($)', 'Flujo acumulado ($)', 'Factor de descuento', 'Flujo descontado ($)',
                  'Acumulado descontado ($)', '¿Recuperado?', '¿Recuperado (descontado)?']


def _hoja_flujo(wb, nombre, refs):
    ws = wb.create_sheet(f'Flujo {nombre}')
    s = fac.ESCENARIOS[nombre]
    col = ESC_COL[nombre]
    _titulo(ws, f'Flujo mensual — escenario {nombre.lower()} (pesos de 2026)',
            'Todas las celdas son fórmulas que leen la hoja «Supuestos» (texto verde: vínculo). Mes 0 = inversión inicial.')
    _encabezados(ws, PRIMERA - 1, COLUMNAS_FLUJO)
    ULT = PRIMERA + s.horizonte_meses
    P = lambda k: f'Supuestos!$C${refs[k]}'
    S = lambda k: f'Supuestos!${col}${refs[k]}'
    obj = '$V$7'
    for i, mes in enumerate(range(s.horizonte_meses + 1)):
        r = PRIMERA + i
        a = f'A{r}'
        ws[a] = mes
        ws[a].font = _f(AZUL)
        ws[f'B{r}'] = (f'=IF({a}=0,0,IF({a}<={P("meses_piloto")},{P("familias_piloto")},IF({a}>={S("meses_estado_estable")},{obj},'
                       f'{P("familias_piloto")}+({obj}-{P("familias_piloto")})*({a}-{P("meses_piloto")})/({S("meses_estado_estable")}-{P("meses_piloto")}))))')
        ws[f'C{r}'] = f'=B{r}*{S("horas_semana_familia")}'
        ws[f'D{r}'] = f'=C{r}*{P("semanas_mes")}'
        ws[f'E{r}'] = f'=IF({a}=0,0,MAX({P("cuidadoras_piloto")},ROUNDUP(C{r}/{P("horas_semana_cuidadora")},0)))'
        if mes == 0:
            ws[f'F{r}'] = 0
        else:
            ws[f'F{r}'] = f'=MAX(0,E{r}-IF(A{r - 1}=0,{P("cuidadoras_piloto")},E{r - 1}))+{P("rotacion_mensual")}*E{r}'
        ws[f'G{r}'] = f'=IF({a}=0,0,MAX(1,ROUNDUP(B{r}/{P("familias_por_trabajadora_social")},0)))'
        ws[f'H{r}'] = f'=D{r}*({P("copago_hora")}+{S("subsidio_hora")})'
        ws[f'I{r}'] = f'=D{r}*{P("costo_variable_hora")}'
        ws[f'J{r}'] = f'=IF({a}=0,0,{P("fijos_base")}+G{r}*{P("trabajadora_social")})'
        ws[f'K{r}'] = f'=F{r}*{P("capacitacion_cuidadora")}'
        ws[f'L{r}'] = f'=IF({a}=0,{P("inversion_inicial")},0)'
        ws[f'M{r}'] = f'=IF({a}=0,-L{r},H{r}-I{r}-J{r}-K{r})'
        ws[f'N{r}'] = f'=M{r}' if mes == 0 else f'=N{r - 1}+M{r}'
        ws[f'O{r}'] = f'=1/(1+{P("tasa_mensual")})^{a}'
        ws[f'P{r}'] = f'=M{r}*O{r}'
        ws[f'Q{r}'] = f'=P{r}' if mes == 0 else f'=Q{r - 1}+P{r}'
        ws[f'R{r}'] = f'=IF(AND({a}>0,N{r}>=0),1,0)'
        ws[f'S{r}'] = f'=IF(AND({a}>0,Q{r}>=0),1,0)'
        for letra, formato in zip('BCDEFGHIJKLMNOPQRS', [NUM1, NUM1, NUM0, NUM0, NUM1, NUM0, PESOS, PESOS, PESOS, PESOS,
                                                          PESOS, PESOS, PESOS, '0.0000', PESOS, PESOS, '0', '0']):
            celda = ws[f'{letra}{r}']
            celda.number_format = formato
            celda.font = _f(VERDE if letra in 'HIJL' else NEGRO)
        if mes % 2 == 0:
            for letra in 'ABCDEFGHIJKLMNOPQRS':
                ws[f'{letra}{r}'].fill = GRIS

    rng = lambda letra: f'{letra}{PRIMERA}:{letra}{ULT}'
    rev = f'({P("mes_revision")}+1)'
    ind = [
        ('Escenario', nombre, None),
        ('Tarifa por hora (copago + subsidio)', f'={P("copago_hora")}+{S("subsidio_hora")}', PESOS2),
        ('Margen por hora', f'=V5-{P("costo_variable_hora")}', PESOS2),
        ('Familias objetivo (estado estable)', f'=MAX({P("hogares_zona")}*{S("adopcion")},{P("familias_piloto")})', NUM1),
        ('VAN a la tasa de descuento', f'=SUM({rng("P")})', PESOS),
        ('TIR mensual', f'=IFERROR(IRR({rng("M")},0.02),"No existe")', '0.00%'),
        ('TIR anual equivalente', '=IFERROR((1+V9)^12-1,"No existe")', PCT),
        ('Payback (meses)', f'=IFERROR(INDEX({rng("A")},MATCH(1,{rng("R")},0)),"No se recupera")', NUM0),
        ('Payback descontado (meses)', f'=IFERROR(INDEX({rng("A")},MATCH(1,{rng("S")},0)),"No se recupera")', NUM0),
        ('Costos fijos en estado estable', f'=J{ULT}', PESOS),
        ('Punto de equilibrio (horas al mes)', '=IF(V6>0,V13/V6,"No existe")', NUM0),
        ('Punto de equilibrio (familias)', f'=IF(V6>0,V14/{P("semanas_mes")}/{S("horas_semana_familia")},"No existe")', NUM1),
        ('Utilidad neta acumulada (60 meses)', f'=N{ULT}', PESOS),
        ('Inversión inicial', f'=L{PRIMERA}', PESOS),
        ('Capital de operación (3 meses)', f'=3*(J{PRIMERA + 1}+I{PRIMERA + 1})', PESOS),
        ('Colchón para imprevistos', f'={P("colchon")}*(V17+V18)', PESOS),
        ('Requerimiento total de capital', '=V17+V18+V19', PESOS),
        ('Déficit acumulado máximo', f'=-MIN(0,MIN({rng("N")}))', PESOS),
        ('Familias en estado estable', f'=B{ULT}', NUM1),
        ('Horas de cuidado al mes en estado estable', f'=D{ULT}', NUM0),
        ('Cuidadoras empleadas en estado estable', f'=E{ULT}', NUM0),
        ('Familias en el mes de revisión', f'=INDEX({rng("B")},{rev})', NUM1),
        ('¿Se cierra en la revisión?', f'=IF(ISNUMBER(V15),IF(V25<{P("umbral_continuar")}*V15,"Sí","No"),"Sí")', None),
        ('Costo de cierre', f'={P("meses_costo_cierre")}*INDEX({rng("J")},{rev})', PESOS),
        ('VAN con revisión en el mes 12', f'=IF(V26="Sí",SUM(P{PRIMERA}:INDEX({rng("P")},{rev}))-V27/(1+{P("tasa_mensual")})^{rev},V8)', PESOS),
        ('Utilidad neta con revisión', f'=IF(V26="Sí",SUM(M{PRIMERA}:INDEX({rng("M")},{rev}))-V27,V16)', PESOS),
    ]
    _encabezados(ws, 3, ['Indicador', 'Valor'], col_inicial=21)
    filas_ind = {}
    for j, (etq, formula, formato) in enumerate(ind):
        r = 4 + j
        _celda(ws, f'U{r}', etq, negrita=etq.startswith(('VAN', 'TIR anual', 'Payback (', 'Punto de equilibrio (h', 'Requerimiento')))
        _celda(ws, f'V{r}', formula, formato=formato)
        filas_ind[etq] = r
    _anchos(ws, {'A': 6, **{get_column_letter(i): 13 for i in range(2, 20)}, 'T': 3, 'U': 42, 'V': 16})
    ws.freeze_panes = f'B{PRIMERA}'
    return ws, filas_ind


# =========================================================================================
# Hoja «Escenarios»
# =========================================================================================

def _hoja_escenarios(wb, refs, filas_ind, subsidio_equilibrio):
    ws = wb.create_sheet('Escenarios')
    _titulo(ws, 'Tabla comparativa de escenarios y valor esperado',
            'Vínculos (verde) a las hojas de flujo. El valor esperado pondera con las probabilidades de «Supuestos».')
    _encabezados(ws, 4, ['Indicador', 'Pesimista', 'Realista', 'Optimista', 'Valor esperado'])
    filas = [('Probabilidad asignada', None, PCT), ('VAN', 'VAN a la tasa de descuento', PESOS),
             ('VAN con revisión en el mes 12', 'VAN con revisión en el mes 12', PESOS),
             ('¿Se cierra en la revisión del mes 12?', '¿Se cierra en la revisión?', None),
             ('TIR anual', 'TIR anual equivalente', PCT), ('Payback (meses)', 'Payback (meses)', NUM0),
             ('Payback descontado (meses)', 'Payback descontado (meses)', NUM0),
             ('Punto de equilibrio (horas al mes)', 'Punto de equilibrio (horas al mes)', NUM0),
             ('Punto de equilibrio (familias)', 'Punto de equilibrio (familias)', NUM1),
             ('Utilidad neta acumulada (60 meses)', 'Utilidad neta acumulada (60 meses)', PESOS),
             ('Utilidad neta con revisión', 'Utilidad neta con revisión', PESOS),
             ('Requerimiento total de capital', 'Requerimiento total de capital', PESOS),
             ('Déficit acumulado máximo', 'Déficit acumulado máximo', PESOS),
             ('Familias en estado estable', 'Familias en estado estable', NUM1),
             ('Horas de cuidado al mes en estado estable', 'Horas de cuidado al mes en estado estable', NUM0),
             ('Cuidadoras empleadas en estado estable', 'Cuidadoras empleadas en estado estable', NUM0)]
    for i, (etq, clave, formato) in enumerate(filas):
        r = 5 + i
        _celda(ws, f'A{r}', etq, negrita=etq.startswith('VAN'))
        for j, nombre in enumerate(ESC_COL):
            letra = 'BCD'[j]
            if clave is None:
                valor = f"=Supuestos!{ESC_COL[nombre]}{refs['probabilidad']}"
            else:
                valor = f"='Flujo {nombre}'!V{filas_ind[clave]}"
            _celda(ws, f'{letra}{r}', valor, VERDE, formato)
        if etq.startswith('VAN') or etq.startswith('Utilidad'):
            _celda(ws, f'E{r}', f'=SUMPRODUCT($B$5:$D$5,B{r}:D{r})', formato=formato, negrita=True)
    r = 5 + len(filas) + 1
    _celda(ws, f'A{r}', 'Subsidio mínimo por hora para VAN ≥ 0 (realista)')
    _celda(ws, f'B{r}', subsidio_equilibrio, AZUL, PESOS2,
           nota='Calculado por bisección en el notebook 03 con backend/observatorio/factibilidad.py; no se recalcula aquí.')
    r += 2
    _celda(ws, f'A{r}', 'Decisión recomendada', negrita=True)
    _celda(ws, f'A{r + 1}', 'INVERTIR CONDICIONADO A: (1) convenio de subsidio ≥ $40 por hora con el Sistema Público de Cuidados; '
                            '(2) revisión en el mes 12 con al menos 50 % de las familias del punto de equilibrio; (3) financiamiento '
                            'comprometido por el déficit acumulado máximo del escenario realista más 10 %.')
    ws[f'A{r + 1}'].alignment = Alignment(wrap_text=True, vertical='top')
    ws.merge_cells(f'A{r + 1}:E{r + 3}')
    ws.row_dimensions[r + 1].height = 30
    _anchos(ws, {'A': 46, 'B': 16, 'C': 16, 'D': 16, 'E': 18})
    return ws


# =========================================================================================
# Hojas de pronóstico, modelo y riesgos
# =========================================================================================

def _hoja_pronostico(wb, semestral, backtest, pronostico):
    ws = wb.create_sheet('Pronóstico oferta')
    _titulo(ws, 'Serie semestral de la oferta de respaldo (DENUE) y pronóstico ingenuo',
            'Datos del DENUE en azul. El backtesting y los intervalos del modelo ingenuo son fórmulas; las métricas de los '
            'otros modelos se calcularon en Python (statsmodels) y se copian como valores.')
    series = ['Total de respaldo', 'Guardería', 'Residencia / asilo', 'Centro de día']
    _encabezados(ws, 4, ['Índice', 'Semestre'] + series + ['Δ a 1 semestre', 'Δ a 2 semestres', 'Δ a 3 semestres'])
    n = len(semestral)
    R0 = 5
    for i, (periodo, fila) in enumerate(semestral.iterrows()):
        r = R0 + i
        _celda(ws, f'A{r}', i, AZUL)
        _celda(ws, f'B{r}', periodo, AZUL)
        for j, s in enumerate(series):
            _celda(ws, f'{"CDEF"[j]}{r}', int(fila[s]), AZUL, NUM0)
        for h, letra in zip((1, 2, 3), 'GHI'):
            if i + h < n:
                _celda(ws, f'{letra}{r}', f'=C{r + h}-C{r}', formato=NUM0)
    ULT = R0 + n - 1
    ws[f'B{ULT + 1}'] = 'Semestres sin edición (2023-S1, 2025-S2) llevan el valor anterior.'
    ws[f'B{ULT + 1}'].font = _f(italica=True, tam=9)

    # Backtesting del ingenuo (origen móvil desde el índice 12)
    r = ULT + 3
    _celda(ws, f'A{r}', 'Backtesting del modelo ingenuo, serie total (origen móvil)', negrita=True)
    r += 1
    _encabezados(ws, r, ['h', 'Origen (índice)', 'Real', 'Pronóstico', 'Error', '|Error| / Real', 'Error²'])
    inicio = {}
    for h in (1, 3):
        inicio[h] = r + 1
        for o in range(12, n - h + 1):
            r += 1
            _celda(ws, f'A{r}', h, AZUL)
            _celda(ws, f'B{r}', o, AZUL)
            _celda(ws, f'C{r}', f'=INDEX($C${R0}:$C${ULT},B{r}+A{r})', formato=NUM0)   # y[o+h-1], índice base 1
            _celda(ws, f'D{r}', f'=INDEX($C${R0}:$C${ULT},B{r})', formato=NUM0)        # y[o-1]
            _celda(ws, f'E{r}', f'=C{r}-D{r}', formato=NUM0)
            _celda(ws, f'F{r}', f'=ABS(E{r})/C{r}', formato=PCT)
            _celda(ws, f'G{r}', f'=E{r}^2', formato=NUM0)
        inicio[h] = (inicio[h], r)
    r += 2
    _celda(ws, f'A{r}', 'Comparación de modelos (origen móvil)', negrita=True)
    r += 1
    _encabezados(ws, r, ['Modelo', 'h', 'MAPE', 'RMSE', 'MAE', 'Evaluaciones', 'Origen del cálculo'])
    for h in (1, 3):
        a, b = inicio[h]
        r += 1
        _celda(ws, f'A{r}', 'Ingenuo (elegido)', negrita=True)
        _celda(ws, f'B{r}', h)
        _celda(ws, f'C{r}', f'=AVERAGE(F{a}:F{b})', formato='0.00%')
        _celda(ws, f'D{r}', f'=SQRT(AVERAGE(G{a}:G{b}))', formato=NUM1)
        _celda(ws, f'E{r}', f'=SUMPRODUCT(ABS(E{a}:E{b}))/COUNT(E{a}:E{b})', formato=NUM1)
        _celda(ws, f'F{r}', f'=COUNT(E{a}:E{b})', formato=NUM0)
        _celda(ws, f'G{r}', 'Fórmula')
        for _, f in backtest[(backtest.h == h) & (backtest.modelo != 'Ingenuo')].iterrows():
            r += 1
            _celda(ws, f'A{r}', f.modelo)
            _celda(ws, f'B{r}', h)
            _celda(ws, f'C{r}', f.MAPE / 100, AZUL, '0.00%')
            _celda(ws, f'D{r}', f.RMSE, AZUL, NUM1)
            _celda(ws, f'E{r}', f.MAE, AZUL, NUM1)
            _celda(ws, f'F{r}', int(f.n), AZUL, NUM0)
            _celda(ws, f'G{r}', 'Python (statsmodels), valor copiado')
    r += 2
    _celda(ws, f'A{r}', 'Pronóstico de la oferta total de respaldo (modelo ingenuo, intervalos empíricos)', negrita=True)
    r += 1
    _encabezados(ws, r, ['Periodo', 'h', 'Pronóstico', 'Lím. inf. 80 %', 'Lím. sup. 80 %', 'Lím. inf. 95 %', 'Lím. sup. 95 %'])
    for h, periodo in zip((1, 2, 3), ['2026-S2', '2027-S1', '2027-S2']):
        r += 1
        col = 'GHI'[h - 1]
        cambios = f'${col}${R0}:${col}${ULT - h}'
        _celda(ws, f'A{r}', periodo)
        _celda(ws, f'B{r}', h)
        _celda(ws, f'C{r}', f'=$C${ULT}', formato=NUM0)
        for k, q in zip('DEFG', (0.10, 0.90, 0.025, 0.975)):
            _celda(ws, f'{k}{r}', f'=$C${ULT}+PERCENTILE({cambios},{q})', formato=NUM0)
    _anchos(ws, {'A': 34, 'B': 12, 'C': 14, 'D': 14, 'E': 16, 'F': 14, 'G': 30, 'H': 14, 'I': 14})
    ws.freeze_panes = 'C5'
    return ws


def _hoja_modelo(wb, coeficientes, validacion, ejemplo, factor):
    ws = wb.create_sheet('Modelo ENUT')
    _titulo(ws, 'Modelo de horas de cuidado no remunerado por mujer de 15 a 59 (regresión lineal ponderada, ENUT 2024)',
            'Coeficientes estimados en el notebook 03 (valores). La calculadora aplica el modelo a los promedios de un AGEB.')
    _encabezados(ws, 4, ['Variable', 'Descripción', 'Coeficiente (h/semana)', 'Construcción con el Censo 2020',
                         'Valor de ejemplo (totales CDMX)', 'Aporte (h/semana)'])
    r = 5
    for _, f in coeficientes.iterrows():
        _celda(ws, f'A{r}', f.variable)
        _celda(ws, f'B{r}', f.descripcion)
        _celda(ws, f'C{r}', float(f.coeficiente), AZUL, NUM3)
        _celda(ws, f'D{r}', f.variable_censo)
        if f.variable == 'constante':
            _celda(ws, f'E{r}', 1, AZUL, NUM3)
        else:
            _celda(ws, f'E{r}', float(ejemplo[f.variable]), AZUL, NUM3, relleno=AMARILLO)
        _celda(ws, f'F{r}', f'=C{r}*E{r}', formato=NUM3)
        r += 1
    ultima = r - 1
    _celda(ws, f'B{r}', 'Predicción del modelo (h/semana por mujer)', negrita=True)
    _celda(ws, f'F{r}', f'=SUM(F5:F{ultima})', formato=NUM3, negrita=True)
    _celda(ws, f'B{r + 1}', 'Factor de calibración a la ENUT CDMX')
    _celda(ws, f'F{r + 1}', factor, AZUL, NUM3, nota='ENUT CDMX observada ÷ predicción con los totales del Censo de la CDMX.')
    _celda(ws, f'B{r + 2}', 'Predicción calibrada (h/semana por mujer)', negrita=True)
    _celda(ws, f'F{r + 2}', f'=MAX(0,F{r}*F{r + 1})', formato=NUM3, negrita=True)
    r += 4
    _celda(ws, f'A{r}', 'Validación', negrita=True)
    r += 1
    _encabezados(ws, r, ['Modelo', 'Tarea', 'Conjunto', 'Métrica', 'Valor', '¿Elegido?'])
    for _, f in validacion[validacion.tarea.str.startswith('Horas')].iterrows():
        r += 1
        _celda(ws, f'A{r}', f.modelo)
        _celda(ws, f'B{r}', f.tarea)
        _celda(ws, f'C{r}', f.conjunto)
        _celda(ws, f'D{r}', f.metrica)
        _celda(ws, f'E{r}', float(f.valor) / 100 if f.metrica == 'MAPE' else float(f.valor), AZUL,
               '0.00%' if f.metrica == 'MAPE' else NUM3)
        _celda(ws, f'F{r}', 'Sí' if f.elegido else 'No')
    _anchos(ws, {'A': 30, 'B': 58, 'C': 30, 'D': 40, 'E': 22, 'F': 16})
    return ws


def _hoja_riesgos(wb):
    ws = wb.create_sheet('Riesgos')
    _titulo(ws, 'Análisis de riesgos del programa')
    _encabezados(ws, 3, ['Tipo', 'Descripción', 'Probabilidad', 'Impacto', 'Mitigación'])
    for i, fila in enumerate(RIESGOS):
        r = 4 + i
        for j, v in enumerate(fila):
            c = ws.cell(row=r, column=j + 1, value=v)
            c.font = _f()
            c.alignment = Alignment(wrap_text=True, vertical='top')
        ws.row_dimensions[r].height = 58
    _anchos(ws, {'A': 12, 'B': 70, 'C': 13, 'D': 11, 'E': 70})
    return ws


def _hoja_leeme(wb):
    ws = wb.active
    ws.title = 'Léeme'
    _titulo(ws, 'Anexo A — Cálculos del trabajo integrador (Analítica para Negocios, PP7, UNRC)')
    lineas = [
        'Observatorio de la Economía del Cuidado — CDMX. Programa evaluado: Red de Cuidado a Domicilio «Manos que Cuidan» (escenario simulado).',
        '',
        'Hojas:',
        '  Supuestos — todos los supuestos con su fuente. Cambia las celdas azules; las amarillas son las variables por escenario.',
        '  Flujo Pesimista / Realista / Optimista — flujo mensual de 60 meses e indicadores (VAN, TIR, payback, punto de equilibrio, capital).',
        '  Escenarios — tabla comparativa y valor esperado del VAN.',
        '  Pronóstico oferta — serie semestral del DENUE 2016-2026, backtesting e intervalos del pronóstico.',
        '  Modelo ENUT — coeficientes, calculadora por AGEB y métricas de validación del modelo de horas de cuidado.',
        '  Riesgos — matriz de riesgos con mitigación.',
        '',
        'Colores: azul = dato o supuesto editable · negro = fórmula · verde = vínculo a otra hoja · relleno amarillo = supuesto clave.',
        'Unidades: pesos constantes de 2026; tasas reales. Negativos entre paréntesis.',
        '',
        'Se genera con notebooks/anexo_excel.py desde el notebook 03; las reglas son las de backend/observatorio/factibilidad.py.',
        'Fuentes principales: INEGI (DENUE 2016-2026, Censo 2020, ENUT 2024, ENASIC 2022); SHCP, Unidad de Inversiones (2022), tasa social',
        'de descuento; CONASAMI, salarios mínimos 2026; Ley del Sistema Público de Cuidados de la CDMX (2026).',
    ]
    for i, t in enumerate(lineas):
        c = ws.cell(row=3 + i, column=1, value=t)
        c.font = _f(negrita=t.endswith(':'))
    for i, (texto, color, relleno) in enumerate([('Supuesto editable', AZUL, None), ('Fórmula', NEGRO, None),
                                                  ('Vínculo a otra hoja', VERDE, None), ('Supuesto clave', AZUL, AMARILLO)]):
        c = ws.cell(row=21 + i, column=1, value=texto)
        c.font = _f(color)
        if relleno:
            c.fill = relleno
    ws.column_dimensions['A'].width = 140


def construir(ruta, *, resultados, oferta, semestral, backtest, pronostico, coeficientes, validacion,
              hogares_zona, subsidio_equilibrio, ejemplo, factor):
    ruta = Path(ruta)
    wb = Workbook()
    _hoja_leeme(wb)
    refs = _hoja_supuestos(wb, hogares_zona)
    filas_ind = None
    for nombre in ESC_COL:
        _, filas_ind = _hoja_flujo(wb, nombre, refs)
    _hoja_escenarios(wb, refs, filas_ind, subsidio_equilibrio)
    _hoja_pronostico(wb, semestral, backtest, pronostico)
    _hoja_modelo(wb, coeficientes, validacion, ejemplo, factor)
    _hoja_riesgos(wb)
    # openpyxl no guarda valores calculados: que Excel recalcule todo al abrir.
    wb.calculation.fullCalcOnLoad = True
    wb.save(ruta)
    return ruta
